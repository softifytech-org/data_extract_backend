"""Excel export - the demo output format for showing this prototype to
stakeholders. Exports the current field values (human-reviewed where
edited) for every record in the document, one row per record, in the
same top-to-bottom page order they were detected in.

Kept as an isolated serializer, mirroring how the XML serializer plugs in
(see services/export/xml_export.py) - neither touches OCR, extraction, or
validation.
"""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from app.models.document import DocumentRecord
from app.services.extraction.fields import FIELD_ORDER, FIELD_SPEC_BY_NAME

HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)
REVIEW_FILL = PatternFill(start_color="FFF3CD", end_color="FFF3CD", fill_type="solid")


def export_record_to_excel(record: DocumentRecord, output_dir: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Records"

    headers = ["Document ID", "Record #", "Page"] + [
        FIELD_SPEC_BY_NAME[name].label for name in FIELD_ORDER
    ] + ["Record Confidence"]
    ws.append(headers)
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT

    review_field_cols: dict[int, list[int]] = {}
    for row_offset, entry in enumerate(record.records, start=2):
        row = [record.document_id, entry.record_index, entry.page]
        needs_review_cols: list[int] = []
        for col_offset, name in enumerate(FIELD_ORDER, start=4):
            f = entry.fields.get(name)
            row.append(f.value if f else "")
            if f and f.status.value == "needs_review":
                needs_review_cols.append(col_offset)
        row.append(entry.overall_confidence)
        ws.append(row)
        review_field_cols[row_offset] = needs_review_cols

    for row_idx, cols in review_field_cols.items():
        for col_idx in cols:
            ws.cell(row=row_idx, column=col_idx).fill = REVIEW_FILL

    for col_idx, header in enumerate(headers, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = max(14, len(str(header)) + 2)
    ws.freeze_panes = "A2"

    output_dir.mkdir(parents=True, exist_ok=True)
    output_path = output_dir / f"{record.document_id}.xlsx"
    wb.save(output_path)
    return output_path
